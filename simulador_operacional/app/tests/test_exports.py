"""Exportacoes de resultados consolidados -- ORION-C-23.

export_resultados_pdf e export_resultados_excel nunca tiveram teste. Os dois
montam a resposta inteira na memoria (reportlab e openpyxl escrevem direto no
HttpResponse, nao em arquivo -- nenhum tempfile e necessario aqui) e filtram
por rodada, por grupo e pelo jogo do usuario logado. O risco sem teste: um
filtro quebrado devolve os dados do grupo errado no relatorio que o
gamemaster baixa para fechar a rodada.
"""

from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from io import BytesIO
from unittest.mock import patch

from app.models import ConsolidadoRodadaGrupo, GameConfig, Grupo, Jogo


def _consolidado(grupo, rodada, **campos):
    valores = {
        "receita_vendas": Decimal("10000"),
        "custos_operacionais": Decimal("4000"),
        "custos_totais": Decimal("5000"),
        "lucro_rodada": Decimal("5000"),
        "saldo_caixa": Decimal("50000"),
        "score_multicriterio": Decimal("70"),
    }
    valores.update(campos)
    return ConsolidadoRodadaGrupo.objects.create(grupo=grupo, rodada=rodada, **valores)


@override_settings(MIGRATION_MODULES={"app": None}, SECURE_SSL_REDIRECT=False)
class ExportResultadosTest(TestCase):
    def setUp(self):
        User = get_user_model()
        self.staff = User.objects.create_user(
            email_usuario="staff-export@example.com",
            nome_usuario="Staff",
            password="pass",
            tipo_usuario="gamemaster",
        )
        self.grupo1 = Grupo.objects.create(nome="Equipe Um")
        self.grupo2 = Grupo.objects.create(nome="Equipe Dois")
        _consolidado(self.grupo1, 1)
        _consolidado(self.grupo2, 2)
        self.client.login(username="staff-export@example.com", password="pass")

    # -- PDF --

    def test_export_pdf_devolve_200_e_content_type_pdf(self):
        response = self.client.get(reverse("export_resultados_pdf"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_export_pdf_filtra_por_rodada(self):
        with patch("app.views.canvas.Canvas") as canvas:
            response = self.client.get(reverse("export_resultados_pdf"), {"rodada": 1})

        self.assertEqual(response.status_code, 200)
        linhas = [chamada.args[2] for chamada in canvas.return_value.drawString.call_args_list]
        self.assertTrue(any("Equipe Um" in linha for linha in linhas))
        self.assertFalse(any("Equipe Dois" in linha for linha in linhas))

    def test_export_pdf_filtra_por_grupo(self):
        with patch("app.views.canvas.Canvas") as canvas:
            response = self.client.get(reverse("export_resultados_pdf"), {"grupo": self.grupo2.id})

        self.assertEqual(response.status_code, 200)
        linhas = [chamada.args[2] for chamada in canvas.return_value.drawString.call_args_list]
        self.assertTrue(any("Equipe Dois" in linha for linha in linhas))
        self.assertFalse(any("Equipe Um" in linha for linha in linhas))

    def test_export_pdf_sem_dados_ainda_devolve_200(self):
        response = self.client.get(reverse("export_resultados_pdf"), {"rodada": 999})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    # -- Excel --

    def test_export_excel_devolve_200_e_content_type_xlsx(self):
        response = self.client.get(reverse("export_resultados_excel"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_export_excel_filtra_por_rodada(self):
        response = self.client.get(reverse("export_resultados_excel"), {"rodada": 1})

        planilha = load_workbook(BytesIO(response.content)).active
        nomes_grupo = [linha[0].value for linha in planilha.iter_rows(min_row=2)]
        self.assertIn("Equipe Um", nomes_grupo)
        self.assertNotIn("Equipe Dois", nomes_grupo)

    def test_export_excel_filtra_por_grupo(self):
        response = self.client.get(reverse("export_resultados_excel"), {"grupo": self.grupo1.id})

        planilha = load_workbook(BytesIO(response.content)).active
        nomes_grupo = [linha[0].value for linha in planilha.iter_rows(min_row=2)]
        self.assertIn("Equipe Um", nomes_grupo)
        self.assertNotIn("Equipe Dois", nomes_grupo)

    def test_export_excel_sem_dados_devolve_so_o_cabecalho(self):
        response = self.client.get(reverse("export_resultados_excel"), {"rodada": 999})

        self.assertEqual(response.status_code, 200)
        planilha = load_workbook(BytesIO(response.content)).active
        self.assertEqual(planilha.max_row, 1)
        self.assertEqual(planilha.cell(row=1, column=1).value, "Grupo")


@override_settings(MIGRATION_MODULES={"app": None}, SECURE_SSL_REDIRECT=False)
class ExportResultadosEscopoPorGrupoTest(TestCase):
    """Usuario comum so pode ver o consolidado do proprio jogo (nao-staff)."""

    def setUp(self):
        User = get_user_model()
        self.membro = User.objects.create_user(
            email_usuario="membro-export@example.com",
            nome_usuario="Membro",
            password="pass",
        )
        self.grupo_do_usuario = Grupo.objects.create(nome="Meu Grupo")
        self.grupo_do_usuario.membros.add(self.membro)
        # Grupo alheio precisa estar em outro jogo: dois grupos sem jogo
        # caem no mesmo balde (jogo__isnull=True) e o filtro nao teria o que
        # provar.
        config = GameConfig.objects.create()
        jogo_alheio = Jogo.objects.create(nome="Outro Jogo", config=config)
        self.outro_grupo = Grupo.objects.create(nome="Grupo Alheio", jogo=jogo_alheio)
        _consolidado(self.grupo_do_usuario, 1)
        _consolidado(self.outro_grupo, 1)
        self.client.login(username="membro-export@example.com", password="pass")

    def test_usuario_sem_grupo_recebe_pdf_sem_dados(self):
        self.membro.grupos.clear()

        response = self.client.get(reverse("export_resultados_pdf"))

        self.assertEqual(response.status_code, 200)
        self.assertNotIn(b"Meu Grupo", response.content)
        self.assertNotIn(b"Grupo Alheio", response.content)

    def test_export_excel_nao_vaza_grupo_de_outro_jogo(self):
        response = self.client.get(reverse("export_resultados_excel"))

        planilha = load_workbook(BytesIO(response.content)).active
        nomes_grupo = [linha[0].value for linha in planilha.iter_rows(min_row=2)]
        self.assertIn("Meu Grupo", nomes_grupo)
        self.assertNotIn("Grupo Alheio", nomes_grupo)
